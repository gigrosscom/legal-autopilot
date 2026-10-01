"""Zann court practice (konsilier/zann/court.py): sud.kz pages in their own markup behind httpx.MockTransport —
the tree of normative resolutions, practice reviews, bulletins — the priority queue, gzip originals and texts in
the storage, the manifest, robots.txt and its Crawl-delay, back-off, imports and the anonymised export. No network."""

from __future__ import annotations

import gzip
import io
import json
import os
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import unquote
from zoneinfo import ZoneInfo

import httpx
import pytest
from docx import Document
from sqlalchemy import select

from konsilier.core.adapters.storage import LocalStorage
from konsilier.core.db import make_engine, make_session_factory
from konsilier.core.models import Base, ZannCourtDoc, ZannCourtPage
from konsilier.zann import court
from konsilier.zann.anonymize import load_rules
from konsilier.zann.court import (MANIFEST_KEY, Collector, CourtFetcher, court_metrics, export_anonymised, file_ext,
                                  load_site)

PACKS = Path(__file__).resolve().parents[3] / "packs"
SITE = load_site(PACKS, "kz")
RULES = load_rules(PACKS, "kz")
classify, priority, date_and_number = SITE.classify, SITE.priority, SITE.date_and_number

NP_LABOUR_RU = ("Нормативное постановление Верховного Суда Республики Казахстан от 15 ноября 2024 года № 5. "
                "О некоторых вопросах применения судами законодательства при разрешении трудовых споров. "
                "1. Судам разъясняется, что трудовой договор заключается в письменной форме. " * 6)
NP_LABOUR_KK = ("Қазақстан Республикасы Жоғарғы Сотының нормативтік қаулысы. Еңбек дауларын шешу кезінде соттардың "
                "заңнаманы қолдануының кейбір мәселелері туралы. Соттарға еңбек шарты жазбаша нысанда жасалатыны "
                "түсіндіріледі. " * 6)
NP_OLD = "О судебной практике по делам о взыскании кредитной задолженности. Банк вправе требовать возврата займа. " * 6
REVIEW = ("Обобщение судебной практики по делам о защите прав потребителей. Истец Ахметов Болат Серикович "
          "(ИИН 850412300123, тел. +7 701 234 56 78) купил товар ненадлежащего качества; суд взыскал с ТОО "
          "«Альфа» стоимость товара. Ахметову возмещен моральный вред. " * 4)


def docx_bytes(text: str) -> bytes:
    d = Document()
    for line in text.split(". "):
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def blank_pdf() -> bytes:
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def shell(title: str, body: str, menu: str = "") -> str:
    """A sud.kz page: a menu (repeats links), the title, the content, the sidebar."""
    return (f'<html><body><div id="menu">{menu}</div><h1 class="title" id="page-title">{title}</h1>'
            f'<div class="region region-help">{body}</div><div class="region region-sidebar-second">'
            f'<a href="https://sud.kz/sites/default/files/pagefiles/sidebar.pdf">не документ</a></div>'
            f'<div id="footer"></div></body></html>')


TREE = ('<div class="html-tree"><ul id="menu" class="menu"><li><div class="content">'
        "<a href='#' class=\"node-link\">Утратившие силу</a></div><ul class=\"menu\"><li><div class=\"content\">"
        "<a href='/rus/legislation/CAT01/79696/2010' class=\"node-link\">2010</a></div></li></ul></li>"
        "<li><div class=\"content\"><a href='#' class=\"node-link\">По гражданскому и гражданско-процессуальному "
        "законодательству</a></div><ul class=\"menu\"><li><div class=\"content\">"
        "<a href='/rus/legislation/CAT01/79693/2024' class=\"node-link\">2024</a></div></li></ul></li></ul></div>")


def np_item(title: str, year: str, files: list[tuple[int, str]]) -> str:
    links = "".join(f'<div class="field-item"><span class="file"><img class="file-icon" src="/x.png">'
                    f'<a href="http://sud.gov.kz/library/download/{i}">{name}</a></div>' for i, name in files)
    return (f'<p class="library-item"><h2 class="main-group-title">{title}</h2><div class="field">'
            f'<span class="field-label">Год:</span>\n\t\t<span class="field-value">{year}</span></div>'
            f'<div class="field"><span class="field-label">Файлы:</span><div class="field-value field-type-file">'
            f'{links}</div></div></p>')


class Site:
    def __init__(self) -> None:
        self.requests: list[str] = []
        self.robots = ("User-agent: *\nCrawl-delay: 10\nDisallow: /search/\nDisallow: /user/login/\n")
        self.busy: dict[str, int] = {}  # path → 503 answers before it works
        bulletin_menu = ('<a href="/rus/content/byulleteni-za-2025-god">2025</a>'
                         '<a href="/rus/content/byulleten-za-2015-god">2015</a>')
        self.pages = {
            "/rus/legislation": shell("Нормативные постановления ВС", "", TREE),
            "/rus/legislation/CAT01/79693/2024": shell("Нормативные постановления ВС", np_item(
                "О некоторых вопросах применения судами законодательства при разрешении трудовых споров", "2024",
                [(188419, "НП трудов КАЗ.docx"), (188420, "НП трудов РУС.DOCX")]), TREE),
            "/rus/legislation/CAT01/79696/2010": shell("Нормативные постановления ВС", np_item(
                "О судебной практике по делам о взыскании кредитной задолженности", "2010", [(1001, "НП кредит.docx")])),
            "/rus/content/obzory-sudebnoy-praktiki": shell("Обзоры судебной практики", (
                '<a href="https://sud.kz/rus/print/297930">Версия для печати</a>'
                '<a href="https://sud.kz/sites/default/files/pagefiles/obobshchenie_potrebiteli.docx">Обобщение '
                'судебной практики по делам о защите прав потребителей за 2023 год</a>'
                '<a href="https://sud.kz/sites/default/files/pagefiles/old.doc">Обзор по уголовным делам</a>'
                '<a href="https://sud.kz/sites/default/files/pagefiles/gone.pdf">Удалённый файл</a>'
                '<a href="https://office.sud.kz/courtActs/doc.pdf">Акт из банка судебных актов</a>'
                '<a href="https://sud.kz/sites/default/files/pagefiles/obobshchenie_potrebiteli.docx">Обобщение '
                '(повтор)</a>'), bulletin_menu),
            "/rus/kategoriya/byulleten-vs": shell("Бюллетень ВС", '<a href="/rus/content/byulleteni-za-2026-god">'
                                                  'Подробнее о Бюллетени за 2026 год</a>', bulletin_menu),
            "/rus/content/byulleteni-za-2026-god": shell("Бюллетени за 2026 год", (
                '<a href="https://sud.kz/sites/default/files/pagefiles/byulleten_no_1-3_2026.pdf">Бюллетень №1-3</a>')),
            "/rus/content/byulleteni-za-2025-god": shell("Бюллетени за 2025 год", ""),
            "/rus/content/byulleten-za-2015-god": shell("Бюллетень за 2015 год", ""),
        }
        self.files = {
            "/library/download/188419": (docx_bytes(NP_LABOUR_KK), "НП трудов КАЗ.docx"),
            "/library/download/188420": (docx_bytes(NP_LABOUR_RU), "НП трудов РУС.DOCX"),
            "/library/download/1001": (docx_bytes(NP_OLD), "НП кредит.docx"),
            "/sites/default/files/pagefiles/obobshchenie_potrebiteli.docx": (docx_bytes(REVIEW), ""),
            "/sites/default/files/pagefiles/old.doc": (b"\xd0\xcf\x11\xe0binary word", ""),
            "/sites/default/files/pagefiles/byulleten_no_1-3_2026.pdf": (blank_pdf(), ""),
        }

    def __call__(self, request: httpx.Request) -> httpx.Response:
        path = unquote(request.url.raw_path.decode())
        self.requests.append(f"{request.url.host}{path}")
        if self.busy.get(path, 0) > 0:
            self.busy[path] -= 1
            return httpx.Response(503, text="busy")
        if path == "/robots.txt":
            return httpx.Response(200, text=self.robots)
        if request.url.host == "sud.kz" and path in self.pages:
            return httpx.Response(200, text=self.pages[path], headers={"Content-Type": "text/html; charset=utf-8"})
        if path in self.files:
            data, name = self.files[path]
            headers = {"Content-Type": "application/octet-stream"}
            if name:
                headers["Content-Disposition"] = "attachment; filename*=utf-8''" + httpx.URL(
                    "http://x/" + name).raw_path.decode()[1:]
            return httpx.Response(200, content=data, headers=headers)
        return httpx.Response(404, text="<html>Страница не найдена</html>")


@pytest.fixture
def db(tmp_path):
    engine = make_engine(os.environ.get("TEST_DATABASE_URL") or f"sqlite:///{tmp_path}/court.db")
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    yield make_session_factory(engine)
    engine.dispose()


def make(db, tmp_path, site: Site, **kw) -> Collector:
    client = httpx.Client(transport=httpx.MockTransport(site), follow_redirects=True)
    sleeps: list[float] = kw.pop("sleeps", [])
    fetch = CourtFetcher(delay=kw.pop("delay", 0), retries=kw.pop("retries", 0), sleep=sleeps.append, client=client,
                         max_bytes=kw.pop("max_bytes", 10_000_000))
    return Collector(db, kw.pop("storage", None) or LocalStorage(tmp_path / "files"), fetch, SITE, RULES, **kw)


# ------------------------------------------------------------------ parsing
def test_parse_tree_leaf_files_and_years():
    site = Site()
    leaves = SITE.parse_index("np", site.pages["/rus/legislation"], "https://sud.kz/rus/legislation")
    assert leaves == [("https://sud.kz/rus/legislation/CAT01/79696/2010", "Утратившие силу / 2010"),
                      ("https://sud.kz/rus/legislation/CAT01/79693/2024",
                       "По гражданскому и гражданско-процессуальному законодательству / 2024")]
    items = SITE.parse_leaf("np", site.pages["/rus/legislation/CAT01/79693/2024"],
                          "https://sud.kz/rus/legislation/CAT01/79693/2024")
    assert [(f.url, f.lang, f.year, f.status) for f in items] == [
        ("https://sud.gov.kz/library/download/188419", "kk", "2024", "in_force"),
        ("https://sud.gov.kz/library/download/188420", "ru", "2024", "in_force")]
    lost = SITE.parse_leaf("np", site.pages["/rus/legislation/CAT01/79696/2010"],
                          "https://sud.kz/rus/legislation/CAT01/79696/2010")
    assert lost[0].status == "lost"
    files = SITE.parse_leaf("review", site.pages["/rus/content/obzory-sudebnoy-praktiki"], "https://sud.kz/rus/content/x")
    assert [f.url.rsplit("/", 1)[-1] for f in files] == ["obobshchenie_potrebiteli.docx", "old.doc", "gone.pdf"]
    assert files[0].year == "2023"  # sidebar and print links are not documents, a repeated link counts once
    years = SITE.parse_index("bulletin", site.pages["/rus/kategoriya/byulleten-vs"], "https://sud.kz/rus/kategoriya/byulleten-vs")
    assert [y for _, y in years] == ["2026", "2025", "2015"]


def test_classify_and_priority_follow_our_scenarios():
    assert classify("О защите прав потребителей") == "consumer"
    assert classify("при разрешении трудовых споров") == "labour"
    assert classify("Еңбек дауларын шешу") == "labour"
    assert classify("по жилищным спорам") == "housing" and classify("договор аренды") == "housing"
    assert classify("о взыскании алиментов") == "alimony" and classify("о расторжении брака") == "alimony"
    assert classify("по делам об административных правонарушениях") == "fines"
    assert classify("Краткий обзор административных дел, рассмотренных в кассационном порядке") == "state"
    assert classify("о взыскании кредитной задолженности") == "credit"
    assert classify("Обзор по уголовным делам") == "other"
    assert priority("consumer", "np") < priority("labour", "np") < priority("credit", "review") < priority("other", "np")
    assert priority("labour", "np", "lost") > priority("credit", "bulletin")
    assert date_and_number("Постановление от 15 ноября 2024 года № 5") == ("2024-11-15", "5")
    assert date_and_number("Без даты", "2010") == ("2010", None)
    assert file_ext("https://sud.gov.kz/library/download/1", "application/octet-stream",
                    "attachment; filename*=utf-8''%D0%9D%D0%9F.DOCX") == "docx"
    assert file_ext("https://x/a.pdf", "", "") == "pdf"


# ------------------------------------------------------------------ the collector
def test_collects_all_open_sources_in_priority_order(db, tmp_path):
    site = Site()
    storage = LocalStorage(tmp_path / "files")
    sleeps: list[float] = []
    c = make(db, tmp_path, site, storage=storage, sleeps=sleeps)
    stats = c.run()
    assert stats.stopped == "idle" and stats.errors == 0
    assert stats.missing == 1  # gone.pdf
    assert stats.notext == 2   # the .doc and the scan-like PDF are stored without text
    assert c.fetch.delay == 10  # robots.txt Crawl-delay applied
    with db() as s:
        docs = {d.url.rsplit("/", 1)[-1]: d for d in s.scalars(select(ZannCourtDoc))}
        pages = s.scalars(select(ZannCourtPage)).all()
    assert {p.state for p in pages} == {"done"} or {p.state for p in pages} == {"done", "error"}
    assert docs["188420"].state == "done" and docs["188420"].lang == "ru" and docs["188420"].category == "labour"
    assert docs["188419"].lang == "kk" and docs["188419"].court.startswith("Верховный Суд")
    assert docs["1001"].status == "lost" and docs["1001"].category == "credit"
    assert docs["obobshchenie_potrebiteli.docx"].category == "consumer"
    assert docs["old.doc"].state == "notext" and docs["old.doc"].key
    assert docs["gone.pdf"].state == "missing"
    # consumers first, then labour, then the rest; lost-force resolutions last
    order = [r.split("/")[-1] for r in site.requests if "/library/" in r or "/pagefiles/" in r]
    assert order[0] == "obobshchenie_potrebiteli.docx" and order[1:3] in (["188419", "188420"], ["188420", "188419"])
    assert order[-1] == "1001"
    # originals and texts gzip in our storage
    d = docs["188420"]
    assert d.key == f"zann/court/orig/{d.id}.docx.gz" and d.text_key == f"zann/court/text/{d.id}.txt.gz"
    assert "трудовой договор" in gzip.decompress(storage.get(d.text_key)).decode()
    assert gzip.decompress(storage.get(d.key))[:2] == b"PK"
    rows = [json.loads(x) for x in gzip.decompress(storage.get(MANIFEST_KEY)).decode().splitlines()]
    assert {r["id"] for r in rows} == {x.id for x in docs.values() if x.key}
    r = next(r for r in rows if r["id"] == d.id)
    assert r["court"].startswith("Верховный") and r["category"] == "labour" and r["date"] == "2024"
    assert len(r["sha256"]) == 64 and r["chars"] > 300 and r["url"] == "https://sud.gov.kz/library/download/188420"
    with db() as s:
        m = court_metrics(s)
    assert m["done"] == 4 and m["notext"] == 2 and m["missing"] == 1 and m["by_category"]["labour"] == 2
    assert m["pages"]["pending"] == 0 and m["bytes"] > 0


def test_texts_are_stored_masked_originals_kept_and_the_acts_bank_is_never_fetched(db, tmp_path):
    site = Site()
    storage = LocalStorage(tmp_path / "files")
    stats = make(db, tmp_path, site, storage=storage).run()
    assert stats.masked > 0
    assert not [r for r in site.requests if r.startswith("office.sud.kz")]
    with db() as s:
        docs = s.scalars(select(ZannCourtDoc)).all()
    assert not [x for x in docs if "office.sud.kz" in x.url]
    d = next(x for x in docs if x.url.endswith("obobshchenie_potrebiteli.docx"))
    text = gzip.decompress(storage.get(d.text_key)).decode()
    assert "850412300123" not in text and "Ахметов" not in text and "+7 701 234 56 78" not in text
    assert "[ИИН1]" in text and "[ФИО1]" in text and "ТОО «Альфа»" in text
    original = Document(io.BytesIO(gzip.decompress(storage.get(d.key))))
    assert "850412300123" in "".join(p.text for p in original.paragraphs)  # the original is kept as it was
    # a pack that seeds a blocked host does not load; a stray queued url on it is never requested
    data = {"hosts": ["sud.kz"], "blocked_hosts": ["office.sud.kz"],
            "sources": {"x": {"parser": "files", "seeds": [{"url": "https://office.sud.kz/courtActs/index.xhtml"}]}}}
    with pytest.raises(ValueError):
        court.Site(data)
    with db() as s:
        s.add(ZannCourtDoc(id="b" * 40, source="review", url="https://office.sud.kz/courtActs/x.pdf",
                           state="pending", priority=0, attempts=0))
        s.commit()
    site.requests.clear()
    make(db, tmp_path, site, storage=storage).run()
    assert not [r for r in site.requests if r.startswith("office.sud.kz")]
    with db() as s:
        assert s.get(ZannCourtDoc, "b" * 40).state == "blocked"


def test_resumes_after_a_limit_and_refreshes_only_new_documents(db, tmp_path):
    site = Site()
    now = [datetime(2026, 10, 1, tzinfo=timezone.utc)]
    c = make(db, tmp_path, site, now=lambda: now[0])
    first = c.run(limit=2)
    assert first.stopped == "limit" and first.docs == 2
    second = c.run()
    assert second.stopped == "idle" and second.docs == 4 and second.pages == 0
    site.requests.clear()
    assert c.run().stopped == "idle" and not [r for r in site.requests if "robots" not in r]
    # a month later the pages are walked again; known files are not downloaded again, a new one is
    site.pages["/rus/content/byulleteni-za-2026-god"] = site.pages["/rus/content/byulleteni-za-2026-god"].replace(
        "</a>", '</a><a href="https://sud.kz/sites/default/files/pagefiles/byulleten_no_4-6_2026.pdf">№4-6</a>', 1)
    site.files["/sites/default/files/pagefiles/byulleten_no_4-6_2026.pdf"] = (blank_pdf(), "")
    now[0] += timedelta(days=31)
    site.requests.clear()
    third = c.run()
    assert third.discovered == 1 and third.docs == 1
    assert [r for r in site.requests if "/pagefiles/" in r or "/library/" in r] == [
        "sud.kz/sites/default/files/pagefiles/byulleten_no_4-6_2026.pdf"]


def test_discover_only_reads_pages_not_files(db, tmp_path):
    site = Site()
    stats = make(db, tmp_path, site).run(discover_only=True)
    assert stats.stopped == "idle" and stats.docs == 0 and stats.discovered == 7
    assert not [r for r in site.requests if "/library/" in r or "/pagefiles/" in r]


def test_robots_closing_the_site_stops_the_run(db, tmp_path):
    site = Site()
    site.robots = "User-agent: *\nDisallow: /\n"
    stats = make(db, tmp_path, site).run()
    assert stats.stopped == "robots" and stats.pages == 0
    assert all(r.endswith("/robots.txt") for r in site.requests)


def test_unreachable_robots_stops_the_run_as_unreachable(db, tmp_path):
    site = Site()
    site.busy["/robots.txt"] = 99
    stats = make(db, tmp_path, site).run()
    assert stats.stopped == "unreachable" and stats.pages == 0 and stats.errors == 1
    assert stats.error.startswith("robots.txt sud.kz: HTTPStatusError") and "\n" not in stats.error


def test_back_off_on_busy_server_and_size_cap(db, tmp_path):
    site = Site()
    site.busy["/rus/legislation"] = 2
    sleeps: list[float] = []
    c = make(db, tmp_path, site, retries=3, sleeps=sleeps, sources=("np",), max_bytes=5000)
    stats = c.run()
    assert stats.pages == 3
    assert sum(1 for s in sleeps if s >= 10) >= 2  # backed off twice, then got the page
    with db() as s:
        docs = s.scalars(select(ZannCourtDoc)).all()
    assert {d.state for d in docs} == {"error"} and all("larger than" in d.error for d in docs)


def test_errors_in_a_row_stop_the_run(db, tmp_path):
    site = Site()
    for p in [*site.files, "/sites/default/files/pagefiles/gone.pdf"]:
        site.busy[p] = 99
    stats = make(db, tmp_path, site, sources=("np", "review", "bulletin")).run()
    assert stats.stopped == "errors" and stats.errors == court.MAX_ERRORS_IN_A_ROW


# ------------------------------------------------------------------ imports and the anonymised export
def test_import_and_export_only_anonymised_texts(db, tmp_path):
    site = Site()
    storage = LocalStorage(tmp_path / "files")
    c = make(db, tmp_path, site, storage=storage)
    c.run()
    act = ("РЕШЕНИЕ. Истец Серикбаев Нурлан Ерланович (ИИН 900101300456), проживающий: г. Астана, ул. Кенесары, "
           "д. 40, кв. 12, обратился с иском к ТОО «Ромашка» о взыскании заработной платы. Суд решил взыскать "
           "с ответчика в пользу Серикбаева Н.Е. 500 000 тенге. Дело № 7194-24-00-2/1234. " * 3)
    did = c.import_file("act1.txt", act.encode("utf-8"), {"court": "Сарыаркинский районный суд г. Астаны",
                                                         "number": "7194-24-00-2/1234", "url": "import:act1"})
    with db() as s:
        d = s.get(ZannCourtDoc, did)
        assert d.source == "acts" and d.category == "labour" and d.state == "done"
    out = list(export_anonymised(db, storage, RULES))
    assert {r["source"] for r in out} == {"np", "review", "acts"}
    blob = json.dumps(out, ensure_ascii=False)
    for secret in ("850412300123", "Ахметов", "+7 701 234 56 78", "900101300456", "Серикбаев", "Кенесары",
                   "7194-24-00-2/1234", "zann/court/orig", "import:act1"):
        assert secret not in blob, secret
    acts = next(r for r in out if r["source"] == "acts")
    assert acts["url"] is None and acts["number"] is None and acts["anonymised"]["ФИО"] >= 2
    assert "[ФИО1]" in acts["text"] and "[ИИН1]" in acts["text"] and "[АДРЕС1]" in acts["text"]
    assert "ТОО «Ромашка»" in acts["text"] and "500 000 тенге" in acts["text"]
    labour = list(export_anonymised(db, storage, RULES, categories=("labour",), sources=("np",)))
    assert len(labour) == 2 and all(r["url"].startswith("https://sud.gov.kz/") for r in labour)


# ------------------------------------------------------------------ the job and the settings
def test_job_runs_nightly_in_its_own_thread(db, tmp_path):
    site = Site()
    started: list = []
    job = court.make_job(lambda: make(db, tmp_path, site), ZoneInfo("Asia/Almaty"), hour=4, minutes=1)
    job.start = started.append
    three = datetime(2026, 10, 1, 22, 0, tzinfo=timezone.utc)  # 03:00 in Almaty (UTC+5 since 2024)
    assert job(None, three) == 0 and not started
    four = three + timedelta(hours=1)
    job(None, four)
    job(None, four + timedelta(minutes=5))  # once a night
    assert len(started) == 1
    started[0]()
    assert job.last is not None and job.last.stopped == "idle" and job.last.docs == 6


def test_settings_and_build(db, tmp_path):
    from konsilier.config import Settings

    s = Settings()
    assert s.zann_court_enabled is False and s.zann_court_sources == "" and s.zann_court_country == ""
    assert court.countries(s.packs_dir) == ["kz"] and court.pick_country(s.packs_dir) == "kz"
    col = court.build_collector(s, db, LocalStorage(tmp_path / "f"))
    assert col.site.court.startswith("Верховный Суд")
    assert col.sources == ("np", "review", "bulletin") and col.fetch.delay >= 0.5
    with pytest.raises(ValueError):
        Collector(db, None, None, SITE, RULES, sources=("office",))


# ------------------------------------------------------------------ what sud.kz lets through, and the status line
class Picky(Site):
    """sud.kz as it answered on 01.10.2026: the connection is dropped without a response when the User-Agent has
    «collector», «spider», «bot» or «httpx» (robots.txt included); the file service of sud.gov.kz (JBoss) answers
    404 to a User-Agent without a platform token such as «(Windows NT 10.0; …)»."""

    def __call__(self, request: httpx.Request) -> httpx.Response:
        ua = request.headers.get("User-Agent", "").lower()
        if any(w in ua for w in ("collector", "spider", "bot", "httpx")):
            self.requests.append(f"dropped {request.url.host}{request.url.path}")
            raise httpx.RemoteProtocolError("Server disconnected without sending a response.", request=request)
        if request.url.host == "sud.gov.kz" and request.url.path.startswith("/library/") and "(windows nt" not in ua:
            return httpx.Response(404, text="<html>JBWEB000065: HTTP Status 404</html>")
        return super().__call__(request)


def picky_collector(db, tmp_path, site, user_agent: str) -> Collector:
    headers = CourtFetcher().client.headers  # what the collector really sends
    headers["User-Agent"] = user_agent
    client = httpx.Client(transport=httpx.MockTransport(site), follow_redirects=True, headers=headers)
    fetch = CourtFetcher(delay=0, retries=1, sleep=lambda s: None, client=client)
    return Collector(db, LocalStorage(tmp_path / "files"), fetch, SITE, RULES)


def test_honest_user_agent_blocked_robot_filter_is_reported_not_bypassed(db, tmp_path):
    """We name ourselves honestly; a site that drops robots by User-Agent stops the run as «unreachable» (shown as
    state=blocked(robots_unreachable)) — the collector never disguises itself as a browser."""
    fetcher = CourtFetcher()
    assert fetcher.client.headers["User-Agent"] == court.USER_AGENT == SITE.user_agent
    assert "Mozilla" not in court.USER_AGENT and "Konsilier.AI" in court.USER_AGENT
    assert fetcher.client.headers["From"] == court.FROM
    site = Picky()
    stats = picky_collector(db, tmp_path, site, court.USER_AGENT).run()
    assert stats.stopped == "unreachable" and stats.pages == 0
    assert court.Site({"sources": {}, "user_agent": "X/1"}).user_agent == "X/1"


def test_job_keeps_its_state_for_the_status_line(db, tmp_path):
    site = Site()
    site.busy["/robots.txt"] = 99
    started: list = []
    job = court.make_job(lambda: make(db, tmp_path, site), ZoneInfo("Asia/Almaty"), hour=-1, minutes=50,
                         session_factory=db)
    job.start = started.append
    with db() as s:
        assert court_metrics(s)["status"]["state"] == "never_run"
    t = datetime(2026, 10, 1, 19, 0, tzinfo=timezone.utc)
    job(None, t)
    with db() as s:
        assert court.court_status(s)["state"] == "running"
    started[0]()
    with db() as s:
        st = court_metrics(s)
    assert st["status"]["state"] == "blocked(robots_unreachable)"
    assert "robots.txt sud.kz" in st["status"]["last_error"] and st["status"]["last_run"]
    nxt = datetime.fromisoformat(st["status"]["next_run"])
    assert timedelta(minutes=14) < nxt - t < timedelta(minutes=16)  # tried again in 15 minutes
    assert st["pages"] == {"done": 0, "pending": 4, "error": 0}  # the status row is not a page
    # the site answers again: the next run collects and the line says «waiting»
    site.busy.clear()
    job(None, t + timedelta(minutes=16))
    with db() as s:
        assert court.court_status(s)["last_error"]  # the last error stays visible while a run is going
    started[1]()
    with db() as s:
        st = court.court_status(s)
    assert st["state"] == "waiting" and st["last_error"] is None and st["run"]["saved"] == 6

    def no_sources():
        raise court.NoSources("set ZANN_COURT_COUNTRY, packs with court sources: []")
    job2 = court.make_job(no_sources, ZoneInfo("Asia/Almaty"), hour=-1, session_factory=db)
    job2.start = started.append
    job2(None, t + timedelta(hours=3))
    started[-1]()
    with db() as s:
        assert court.court_status(s)["state"] == "no_sources"
    assert court.run_state("failed", "OperationalError: db") == "error(OperationalError)"
    assert court.run_state("robots") == "blocked(robots_closed)"
    assert court.run_state("errors") == "error(errors_in_a_row)"


def test_no_sources_is_its_own_error(tmp_path):
    with pytest.raises(court.NoSources):
        court.pick_country(tmp_path)  # no pack has court sources
    with pytest.raises(court.NoSources):
        court.load_site(tmp_path, "uz")
    with pytest.raises(court.NoSources):
        Collector(None, None, None, SITE, RULES, sources=("nope",))
