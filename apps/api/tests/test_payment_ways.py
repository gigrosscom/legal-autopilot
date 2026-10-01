"""Ways to pay (PAYMENT_METHODS, docs/kaspi-pay-plan.md): Kaspi Pay link, printed Kaspi QR, a Kaspi bill to the
client's number, «Счёт на оплату» for a company. Each ends like the transfer: the clients desk confirms in /ops and
the document becomes downloadable. Off by default: without PAYMENT_METHODS the payment screen is as before."""

from __future__ import annotations

import hashlib
import hmac
import io
import json
from decimal import Decimal

from konsilier.config import Settings
from konsilier.core.adapters.payment import (ManualTransferPaymentAdapter, Requisites, build_payments, kz_phone,
                                             valid_bin)
from konsilier.core.bill import BillWords, amount_in_words, money, number_in_words
from konsilier.identity.senders import LogSender

from .test_e2e import web_user
from .test_lawyer_onboarding import Outbox
from .test_ops_centre import operator
from .test_payment import PHONE, RECIPIENT, confirm, qualified_case

# placeholders only: real requisites live in the server's .env
REQ = Requisites(name="ТОО «Тест»", bin="000000000000", address="Алматы, ул. Тестовая 1", bank="АО «Тест Банк»",
                 iik="KZ00000TEST0000000000", bik="TESTKZKA", director="Директор Тестов Т.")
LINK, QR = "https://pay.kaspi.kz/pay/test", "https://konsilier.com/kaspi-qr-test.png"
DESK = "support@konsilier.com"


def ways(ctx, methods: str = "kaspi_link,kaspi_qr,kaspi_invoice,bank_invoice", requisites: Requisites = REQ):
    ctx.container.engine.payments = ManualTransferPaymentAdapter(
        recipient_name=RECIPIENT, kaspi_phone=PHONE, comment_prefix="ka", ways=methods, kaspi_pay_link=LINK,
        kaspi_qr_image=QR, requisites=requisites)
    ctx.container.engine.config.approval_required_first_n = 0
    ctx.container.sms_sender = LogSender("sms")
    ctx.container.settings.ops_clients_emails = DESK
    outbox = Outbox()
    ctx.container.email_sender = outbox
    return outbox


def open_bill(ctx, purpose: str = "document", phone: str = "+7 701 555 00 11"):
    api, cid = qualified_case(ctx)
    confirm(api, "phone", phone)
    pay = api.post(f"/v1/cases/{cid}/payment", json={"purpose": purpose})["case"]["payment"]
    return api, cid, pay


def desk_confirms(ctx, code: str) -> dict:
    cl = getattr(ctx, "_desk", None) or operator(ctx, DESK)  # the desk signs in once per test
    ctx._desk = cl
    row = next(r for r in cl.get("/v1/ops/clients/payments").json() if r["code"] == code)
    return cl.post(f"/v1/ops/clients/payments/{row['id']}", json={"decision": "paid"})


def document_downloadable(api, cid: str) -> None:
    case = api.post(f"/v1/cases/{cid}/actions/next")["case"]
    a = case["actions"][0]
    assert a["downloadable"] is True
    api.get(f"/v1/cases/{cid}/actions/{a['id']}/document?format=docx")


def test_off_by_default_and_ways_need_their_settings():
    plain = build_payments(Settings(payment_mode="manual_transfer", payment_recipient_name=RECIPIENT,
                                    payment_kaspi_phone=PHONE))
    assert plain.details() == {"recipient_name": RECIPIENT, "kaspi_phone": PHONE}  # as before
    assert not plain.way_available("bank_invoice") and not plain.way_available("kaspi_link")
    # switched on, but the link / QR / requisites are not filled in: only the Kaspi bill (needs nothing) shows
    half = build_payments(Settings(payment_mode="manual_transfer", payment_recipient_name=RECIPIENT,
                                   payment_kaspi_phone=PHONE, payment_methods="kaspi_link, kaspi_qr,bank_invoice,"
                                                                              "kaspi_invoice, nonsense"))
    assert [w["id"] for w in half.details()["ways"]] == ["kaspi_transfer", "kaspi_invoice"]
    full = build_payments(Settings(
        payment_mode="manual_transfer", payment_methods="kaspi_link,kaspi_qr,bank_invoice",
        payment_kaspi_pay_link=LINK, payment_kaspi_qr_image=QR, payment_llp_name=REQ.name, payment_llp_bin=REQ.bin,
        payment_llp_bank=REQ.bank, payment_llp_iik="KZ00 000TEST", payment_llp_bik=REQ.bik))
    # no Kaspi Gold number at all: the other ways still make payment available
    assert full.available() and "kaspi_phone" not in full.details()
    assert full.details()["ways"] == [{"id": "kaspi_link", "url": LINK}, {"id": "kaspi_qr", "image": QR, "url": LINK},
                                      {"id": "bank_invoice", "seller": REQ.name}]
    assert full.requisites.iik == "KZ00000TEST" and full.requisites.kbe == "17" and full.requisites.knp == "859"


def test_helpers():
    assert kz_phone("8 (701) 555-00-11") == "+77015550011" == kz_phone("+7 701 555 00 11") == kz_phone("7015550011")
    assert kz_phone("12345") is None and kz_phone("+1 212 555 0101") is None
    assert valid_bin("0000 0000 0000") == "000000000000" and valid_bin("123") is None
    assert number_in_words(1990) == "одна тысяча девятьсот девяносто"
    assert number_in_words(9990) == "девять тысяч девятьсот девяносто"
    assert number_in_words(2_021_002) == "два миллиона двадцать одна тысяча два"
    assert number_in_words(11_000) == "одиннадцать тысяч"
    tenge = BillWords(currency_forms=("тенге", "тенге", "тенге"), minor="тиын")
    assert amount_in_words(Decimal("29990"), tenge) == "Двадцать девять тысяч девятьсот девяносто тенге 00 тиын"
    assert money(Decimal("1990")) == "1 990,00"


def test_kaspi_link_and_qr_paid_then_document(ctx):
    """Link / QR: the person pays in Kaspi.kz with the code in the message, presses «Оплатить», the desk confirms."""
    outbox = ways(ctx)
    for i, way in enumerate(("kaspi_link", "kaspi_qr")):
        api, cid, pay = open_bill(ctx, phone=f"+7 701 555 00 2{i}")
        assert [w["id"] for w in pay["ways"]] == ["kaspi_transfer", "kaspi_link", "kaspi_qr", "kaspi_invoice",
                                                  "bank_invoice"]
        assert pay["invoice_id"] and pay["way"] is None and pay["kaspi_phone"] == PHONE
        chosen = api.post(f"/v1/invoices/{pay['invoice_id']}/way", json={"way": way})["case"]["payment"]
        assert chosen["way"] == way and chosen["status"] == "pending"
        claimed = api.post(f"/v1/cases/{cid}/payment/claim")["case"]["payment"]
        # the Kaspi Pay link is given on trust (owner 01.10, see below); the QR waits for the desk
        assert claimed["status"] == ("paid" if way == "kaspi_link" else "awaiting_confirmation")
        assert pay["code"] in outbox.sent[-1][2] and "Kaspi" in outbox.sent[-1][2]
        if way == "kaspi_qr":
            assert api.post(f"/v1/cases/{cid}/actions/next")["action_id"] is None  # not before the desk
        assert desk_confirms(ctx, pay["code"])["way"] == way
        document_downloadable(api, cid)


def test_kaspi_bill_to_phone(ctx):
    """The person gives a Kaspi number: the desk is asked at once to send a Kaspi bill; once paid, the document."""
    outbox = ways(ctx)
    api, cid, pay = open_bill(ctx)
    r = api.ctx.client.post(f"/v1/invoices/{pay['invoice_id']}/way", headers=api.h,
                            json={"way": "kaspi_invoice", "phone": "12"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "invalid_phone"
    got = api.post(f"/v1/invoices/{pay['invoice_id']}/way",
                   json={"way": "kaspi_invoice", "phone": "8 701 555 00 11"})["case"]["payment"]
    assert got["status"] == "awaiting_confirmation" and got["payer_phone"] == "+77015550011"
    to, subject, text = outbox.sent[-1][:3]
    assert to == DESK and "выставьте счёт" in subject and "+77015550011" in text and pay["code"] in text
    # while the desk is on it, the person cannot switch to another way
    r = api.ctx.client.post(f"/v1/invoices/{pay['invoice_id']}/way", headers=api.h, json={"way": "kaspi_link"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "invoice_awaiting_confirmation"
    row = desk_confirms(ctx, pay["code"])
    assert row["payer_phone"] == "+77015550011" and row["status"] == "paid"
    document_downloadable(api, cid)


def test_bank_invoice_for_a_company(ctx):
    """«Счёт на оплату»: the buyer's name and БИН → the bill with the company's requisites → transfer → desk."""
    from docx import Document

    ways(ctx)
    api, cid, pay = open_bill(ctx, "case")
    client = api.ctx.client
    iid = pay["invoice_id"]
    assert client.get(f"/v1/invoices/{iid}/bill", headers=api.h).status_code == 409  # not chosen yet
    r = client.post(f"/v1/invoices/{iid}/way", headers=api.h,
                    json={"way": "bank_invoice", "buyer_name": "ТОО «Покупатель»", "buyer_bin": "12"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "invalid_bin"
    got = api.post(f"/v1/invoices/{iid}/way", json={"way": "bank_invoice", "buyer_name": "ТОО «Покупатель»",
                                                    "buyer_bin": "1111 2222 3333"})["case"]["payment"]
    assert got["buyer"] == {"name": "ТОО «Покупатель»", "bin": "111122223333", "address": None}
    assert got["status"] == "pending"
    r = client.get(f"/v1/invoices/{iid}/bill?format=docx", headers=api.h)
    assert r.status_code == 200 and "schet-" in r.headers["content-disposition"]
    doc = Document(io.BytesIO(r.content))
    text = "\n".join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for row in t.rows for c in row.cells])
    for want in (f"Счёт на оплату № {iid}", REQ.name, REQ.bin, REQ.iik, REQ.bik, REQ.bank, "859", "17",
                 "ТОО «Покупатель»", "111122223333", "9 990,00", "Девять тысяч девятьсот девяносто тенге 00 тиын",
                 "Без НДС", pay["code"], "Дело под ключ"):
        assert want in text, want
    assert client.get(f"/v1/invoices/{iid}/bill", headers=api.h).status_code == 200  # PDF, or Word without LibreOffice
    # someone else's bill is not there
    other = web_user(ctx)
    assert client.get(f"/v1/invoices/{iid}/bill", headers=other.h).status_code == 404
    assert client.post(f"/v1/invoices/{iid}/way", headers=other.h, json={"way": "kaspi_link"}).status_code == 404

    api.post(f"/v1/cases/{cid}/payment/claim")
    desk_confirms(ctx, pay["code"])
    view = api.get(f"/v1/cases/{cid}").json()["payment"]
    assert view["case_paid"] is True and view["status"] == "paid"
    document_downloadable(api, cid)


def test_way_that_is_off_is_refused(ctx):
    ways(ctx, methods="kaspi_link", requisites=Requisites())
    api, cid, pay = open_bill(ctx)
    r = api.ctx.client.post(f"/v1/invoices/{pay['invoice_id']}/way", headers=api.h,
                            json={"way": "bank_invoice", "buyer_name": "ТОО «X»", "buyer_bin": "111122223333"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "way_unavailable"


def test_receipt_letter_after_confirmation(ctx):
    outbox = ways(ctx)
    st = ctx.container.settings
    st.payment_receipt_email, st.payment_kaspi_kassa, st.payment_llp_name = True, True, REQ.name
    try:
        api, cid, pay = open_bill(ctx)
        with ctx.container.session_factory() as s:
            from konsilier.core.models import Case, User
            import uuid

            owner = s.get(User, s.get(Case, uuid.UUID(cid)).owner_id)
            owner.email = "client@mail.kz"
            s.commit()
        api.post(f"/v1/invoices/{pay['invoice_id']}/way", json={"way": "kaspi_qr"})
        api.post(f"/v1/cases/{cid}/payment/claim")
        desk_confirms(ctx, pay["code"])
        to, _, text = next(m[:3] for m in reversed(outbox.sent) if m[0] == "client@mail.kz")
        for want in ("Оплата получена", "Сумма: 1 990,00 ₸", "Дата: ", "За что: Подготовка юридического документа",
                     f"Код платежа: {pay['code']}", "Способ: Kaspi QR", f"Продавец: {REQ.name}", f"/case/{cid}",
                     "не фискальный чек", "Фискальный чек за оплату через Kaspi Pay приходит в приложение Kaspi.kz"):
            assert want in text, want
    finally:
        st.payment_receipt_email, st.payment_kaspi_kassa, st.payment_llp_name = False, False, ""


def test_kaspi_webhook_is_off_until_switched_on_and_checks_the_signature(ctx):
    ways(ctx)
    st = ctx.container.settings
    client = ctx.client
    api, cid, pay = open_bill(ctx)
    body = json.dumps({"code": pay["code"], "status": "paid", "amount": 1990, "txn_id": "T-1"}).encode()
    assert client.post("/v1/payments/kaspi/webhook", content=body).status_code == 404  # off
    st.payment_kaspi_webhook, st.payment_kaspi_webhook_secret = True, "test-secret"
    try:
        def sign(raw: bytes) -> dict:
            return {"X-Konsilier-Signature": hmac.new(b"test-secret", raw, hashlib.sha256).hexdigest(),
                    "Content-Type": "application/json"}

        assert client.post("/v1/payments/kaspi/webhook", content=body,
                           headers={"X-Konsilier-Signature": "bad"}).status_code == 401
        # a wrong amount is not «paid»: it goes to the desk
        wrong = json.dumps({"code": pay["code"], "status": "paid", "amount": 990, "txn_id": "T-0"}).encode()
        assert client.post("/v1/payments/kaspi/webhook", content=wrong, headers=sign(wrong)).json()["status"] == \
            "awaiting_confirmation"
        r = client.post("/v1/payments/kaspi/webhook", content=body, headers=sign(body))
        assert r.status_code == 200 and r.json() == {"ok": True, "status": "paid"}
        document_downloadable(api, cid)
        # repeated delivery is harmless
        assert client.post("/v1/payments/kaspi/webhook", content=body, headers=sign(body)).json()["status"] == "paid"
    finally:
        st.payment_kaspi_webhook, st.payment_kaspi_webhook_secret = False, ""


def test_desk_sees_when_kaspi_was_opened(ctx):
    """Owner 01.10 «увидел — нажал — оплатил»: no payment code to type — the desk matches a Kaspi Pay payment by the
    amount and the time of «Оплатить в Kaspi» (the way recorded with its time), shown with the bill in /ops."""
    ways(ctx, methods="kaspi_link,bank_invoice")
    api, cid, pay = open_bill(ctx, phone="+7 701 555 00 31")
    assert [w["id"] for w in pay["ways"]] == ["kaspi_transfer", "kaspi_link", "bank_invoice"]
    cl = getattr(ctx, "_desk", None) or operator(ctx, DESK)
    ctx._desk = cl
    row = next(r for r in cl.get("/v1/ops/clients/payments?status=").json() if r["code"] == pay["code"])
    assert row["kaspi_opened_at"] is None
    api.post(f"/v1/invoices/{pay['invoice_id']}/way", json={"way": "kaspi_link"})
    api.post(f"/v1/cases/{cid}/payment/claim")
    row = next(r for r in cl.get("/v1/ops/clients/payments").json() if r["code"] == pay["code"])
    assert row["kaspi_opened_at"] and row["way"] == "kaspi_link"


def test_kaspi_link_gives_the_document_at_once_and_the_desk_matches_after(ctx):
    """Owner 01.10: «Оплатить» with the Kaspi Pay link = the bill goes to the desk and the document is given at once;
    the desk confirms (nothing given twice) — or does not find it: the person owes it, no new document until paid."""
    outbox = ways(ctx, methods="kaspi_link,bank_invoice")
    api, cid, pay = open_bill(ctx, phone="+7 701 555 00 41")
    api.post(f"/v1/invoices/{pay['invoice_id']}/way", json={"way": "kaspi_link"})
    claimed = api.post(f"/v1/cases/{cid}/payment/claim")["case"]["payment"]
    assert claimed["status"] == "paid" and claimed["trusted"] is True  # the next document can be made now
    assert pay["code"] in outbox.sent[-1][2]  # the desk is told, as before
    document_downloadable(api, cid)
    cl = getattr(ctx, "_desk", None) or operator(ctx, DESK)
    ctx._desk = cl
    row = next(r for r in cl.get("/v1/ops/clients/payments").json() if r["code"] == pay["code"])
    assert row["trusted_at"] and row["status"] == "awaiting_confirmation"
    assert desk_confirms(ctx, pay["code"])["status"] == "paid"
    case = api.get(f"/v1/cases/{cid}").json()
    assert case["payment"]["credits"] == 0 and case["payment"]["owed"] is False  # no second credit


def test_kaspi_link_not_found_means_a_debt(ctx):
    ways(ctx, methods="kaspi_link,bank_invoice")
    api, cid, pay = open_bill(ctx, phone="+7 701 555 00 42")
    api.post(f"/v1/invoices/{pay['invoice_id']}/way", json={"way": "kaspi_link"})
    api.post(f"/v1/cases/{cid}/payment/claim")
    document_downloadable(api, cid)
    cl = getattr(ctx, "_desk", None) or operator(ctx, DESK)
    ctx._desk = cl
    row = next(r for r in cl.get("/v1/ops/clients/payments").json() if r["code"] == pay["code"])
    assert cl.post(f"/v1/ops/clients/payments/{row['id']}", json={"decision": "not_found"})["status"] == "not_found"
    case = api.get(f"/v1/cases/{cid}").json()
    assert case["payment"]["owed"] is True and case["payment"]["status"] == "not_found"
    assert "не нашли вашу оплату" in str(api.get("/v1/notifications").json())  # the reminder, no payment code
    # paying again is a plain claim now: nothing on trust while the person owes
    again = api.post(f"/v1/cases/{cid}/payment/claim")["case"]["payment"]
    assert again["status"] == "awaiting_confirmation"
    assert desk_confirms(ctx, pay["code"])["status"] == "paid"
    assert api.get(f"/v1/cases/{cid}").json()["payment"]["owed"] is False


def test_trust_can_be_switched_off(ctx):
    ways(ctx, methods="kaspi_link")
    ctx.container.engine.config.trust_kaspi_link = False
    api, cid, pay = open_bill(ctx, phone="+7 701 555 00 43")
    api.post(f"/v1/invoices/{pay['invoice_id']}/way", json={"way": "kaspi_link"})
    claimed = api.post(f"/v1/cases/{cid}/payment/claim")["case"]["payment"]
    assert claimed["status"] == "awaiting_confirmation" and claimed["trusted"] is False
