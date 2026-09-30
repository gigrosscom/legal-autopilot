"""«Юрист по кнопке» — the closed pilot (owner's decision 30.09.2026).

The owner marks a few verified lawyers (ЭЦП-tied applications) as pilot lawyers and sets each one's price. The
client sends a request to one of them → the lawyer accepts or declines → after acceptance the client pays the
lawyer's price through the platform, to the COMPANY's account only (the ТОО's Kaspi Pay link or its bank requisites,
never the personal Kaspi Gold used for documents) → once the payment is confirmed the lawyer is assigned to the case
(the pack's consent and engagement papers are generated, signed with ЭЦП as before) and gets the case dossier.

The platform keeps ``lawyer_commission_pct`` of the price (15 % by default); payouts to lawyers are manual in the
pilot. A lawyer bill never marks the case paid and never unlocks documents.
"""

from __future__ import annotations

import base64
import hashlib
import hmac
import io
from datetime import date, datetime, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal
from typing import TYPE_CHECKING, Any

from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session

from .adapters.payment import new_payment_code
from .models import Agreement, Case, Identity, Invoice, LawyerApplication, LawyerRequest, User

if TYPE_CHECKING:
    from .engine import CaseEngine

PURPOSE = "lawyer"
METHOD = "company"  # paid to the company's account; confirmed by the desk / owner (or the Kaspi webhook)


class PilotError(Exception):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def money(value: Decimal | float | int) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def commission(amount: Decimal, pct: float) -> Decimal:
    return money(Decimal(amount) * Decimal(str(pct)) / 100)


# ------------------------------------------------------------------ the company's payment channel
def channel(engine: "CaseEngine") -> dict[str, Any] | None:
    """Where the client pays a lawyer: the ТОО's Kaspi Pay link and / or its requisites. None while neither is set —
    then lawyer payment is unavailable (never a fallback to the personal Kaspi Gold)."""
    cfg = engine.config
    link = (cfg.lawyer_pay_link or "").strip()
    account = (cfg.lawyer_pay_account or "").strip()
    out: dict[str, Any] = {}
    if link.startswith("https://"):
        out["kaspi_pay_link"] = link
    if account:
        out["company_account"] = account
    if not out:
        return None
    out["company_name"] = (cfg.company_name or "").strip() or None
    return out


# ------------------------------------------------------------------ who is in the pilot
INVITE_DAYS = 14


def _invite_sig(secret: str, exp: int) -> str:
    mac = hmac.new(secret.encode(), f"lawyer-invite:{exp}".encode(), hashlib.sha256).digest()
    return base64.urlsafe_b64encode(mac[:16]).decode().rstrip("=")


def invite_token(secret: str, now: datetime, days: int = INVITE_DAYS) -> tuple[str, datetime]:
    """The owner's private link for pilot lawyers: expiry + signature, nothing stored. (token, expires_at)"""
    exp = int((now + timedelta(days=days)).timestamp())
    return f"{exp}.{_invite_sig(secret, exp)}", datetime.fromtimestamp(exp, timezone.utc)


def invite_valid(secret: str, token: str | None, now: datetime) -> bool:
    exp_s, _, sig = (token or "").partition(".")
    if not exp_s.isdigit() or not sig or int(exp_s) < now.timestamp():
        return False
    return hmac.compare_digest(sig, _invite_sig(secret, int(exp_s)))


def is_pilot_lawyer(app: LawyerApplication | None) -> bool:
    return (app is not None and app.status == "verified" and bool(app.pilot) and app.price is not None
            and app.price > 0 and bool(app.iin_hash) and app.user_id is not None)


def pilot_lawyers(session: Session, country: str | None) -> list[LawyerApplication]:
    q = select(LawyerApplication).where(LawyerApplication.status == "verified", LawyerApplication.pilot.is_(True),
                                        LawyerApplication.price.is_not(None)).order_by(LawyerApplication.price,
                                                                                       LawyerApplication.id)
    if country:
        q = q.where(LawyerApplication.country == country.upper())
    return [a for a in session.scalars(q).all() if is_pilot_lawyer(a)]


def lawyer_name(app: LawyerApplication) -> str:
    return app.ecp_name or app.full_name


# ------------------------------------------------------------------ customer–lawyer papers
def iin_identity(session: Session, user_id: Any) -> Identity | None:
    return session.scalar(select(Identity).where(Identity.user_id == user_id, Identity.kind == "iin"))


def render_agreement(pack: Any, kind: str, lang: str, values: dict[str, str]) -> bytes:
    """DOCX from the pack's template. Unreviewed templates carry a visible note; nothing is invented here."""
    ags = pack.agreements
    tpl = ags.templates[kind]
    lang = lang if lang in tpl.body else pack.manifest.default_language
    doc = Document()
    if ags.reviewed_at is None and "unreviewed" in ags.footer:
        doc.add_paragraph(pack.localized(ags.footer["unreviewed"], lang)).runs[0].italic = True
    doc.add_heading(pack.localized(tpl.title, lang), level=1)
    for para in tpl.body[lang]:
        doc.add_paragraph(para.format(**values))
    doc.add_paragraph(values["date"])
    if "signed_with" in ags.footer:
        doc.add_paragraph(pack.localized(ags.footer["signed_with"], lang)).runs[0].italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_agreements(session: Session, engine: "CaseEngine", case: Case, app: LawyerApplication) -> list[Agreement]:
    pack = engine.pack_of(case)
    if pack.agreements is None:
        return []
    lang = pack.lang(case.language)
    owner = session.get(User, case.owner_id)
    ident = iin_identity(session, owner.id)
    sc_title = ""
    if case.scenario_id:
        try:
            sc_title = pack.localized(engine.packs.scenario(case.scenario_id).title, lang)
        except KeyError:
            sc_title = ""
    values = {
        "customer_name": owner.display_name or (case.facts or {}).get("applicant_name") or "—",
        "customer_id": ident.display if ident else "—",
        "lawyer_name": lawyer_name(app),
        "lawyer_kind": pack.localized(pack.agreements.lawyer_kinds.get(app.kind, {"ru": app.kind}), lang),
        "case_ref": str(case.id)[:8].upper(),
        "case_title": sc_title or "—",
        "date": date.today().strftime("%d.%m.%Y"),
    }
    out = []
    for kind in pack.agreements.templates:
        data = render_agreement(pack, kind, lang, values)
        ag = Agreement(case_id=case.id, kind=kind, lawyer_application_id=app.id, docx_key="",
                       sha256=hashlib.sha256(data).hexdigest(),
                       template_reviewed=pack.agreements.reviewed_at is not None)
        session.add(ag)
        session.flush()
        ag.docx_key = engine.storage.put(f"cases/{case.id}/agreements/{ag.id}.docx", data,
                                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        out.append(ag)
    return out


def assign(session: Session, engine: "CaseEngine", case: Case, app: LawyerApplication, actor: str) -> list[Agreement]:
    """Put the lawyer on the case: the cabinet shows it, the papers are generated for both to sign with ЭЦП."""
    if app.status != "verified" or app.user_id is None or not app.iin_hash:
        raise PilotError("lawyer_not_verified")
    if case.lawyer_application_id == app.id:
        raise PilotError("already_assigned")
    case.lawyer_user_id, case.lawyer_application_id = app.user_id, app.id
    ags = create_agreements(session, engine, case, app)
    engine.audit(session, case, actor, "lawyer_assigned", application_id=app.id, agreements=[a.kind for a in ags])
    return ags


# ------------------------------------------------------------------ the lawyer's answer
def decide(session: Session, engine: "CaseEngine", req: LawyerRequest, accept: bool, actor: str) -> None:
    if req.status != "new" or req.application_id is None:
        raise PilotError("request_closed")
    req.status, req.decided_at = ("accepted" if accept else "declined"), utcnow()
    case = session.get(Case, req.case_id)
    app = session.get(LawyerApplication, req.application_id)
    pack = engine.pack_of(case)
    lang = pack.lang(case.language)
    engine.audit(session, case, actor, "lawyer_request_accepted" if accept else "lawyer_request_declined",
                 request=req.id, application_id=req.application_id)
    if accept:
        text = pack.t(lang, "notifications.lawyer_accepted", name=lawyer_name(app),
                      default=f"Юрист {lawyer_name(app)} принял запрос. Оплатите, чтобы он получил материалы дела.")
    else:
        text = pack.t(lang, "notifications.lawyer_declined", name=lawyer_name(app),
                      default=f"Юрист {lawyer_name(app)} не может взять это дело. Выберите другого юриста в "
                              f"карточке дела.")
    engine.notifier.notify(session, case, "lawyer", text)


# ------------------------------------------------------------------ the bill
def open_request(session: Session, case: Case) -> LawyerRequest | None:
    """The case's latest pilot request that is not declined (new, accepted or paid)."""
    return session.scalar(select(LawyerRequest).where(
        LawyerRequest.case_id == case.id, LawyerRequest.application_id.is_not(None),
        LawyerRequest.status.in_(("new", "accepted", "paid"))).order_by(LawyerRequest.id.desc()).limit(1))


def invoice_of(session: Session, req: LawyerRequest) -> Invoice | None:
    if req.invoice_id is None:
        return None
    return session.get(Invoice, req.invoice_id)


def create_invoice(session: Session, engine: "CaseEngine", req: LawyerRequest, actor: str) -> Invoice:
    """The client's bill for the lawyer's price, after the lawyer accepted. The open bill is reused."""
    if req.status == "paid":
        raise PilotError("already_paid")
    if req.status != "accepted" or req.price is None:
        raise PilotError("not_accepted")
    inv = invoice_of(session, req)
    if inv is not None and inv.status in engine.OPEN:
        return inv
    if channel(engine) is None:
        raise PilotError("lawyer_payment_unavailable")
    case = session.get(Case, req.case_id)
    pack = engine.pack_of(case)
    amount = money(req.price)
    pct = Decimal(str(engine.config.lawyer_commission_pct))
    inv = Invoice(case_id=case.id, user_id=case.owner_id, purpose=PURPOSE, code=new_payment_code(), method=METHOD,
                  amount=amount, currency=pack.currency, status="pending", lawyer_request_id=req.id,
                  commission_pct=pct, commission_amount=commission(amount, float(pct)))
    session.add(inv)
    session.flush()
    req.invoice_id = inv.id
    engine.audit(session, case, actor, "invoice_created", invoice=inv.code, purpose=PURPOSE, status=inv.status,
                 amount=str(inv.amount), currency=inv.currency, commission=str(inv.commission_amount))
    return inv


def apply_paid(session: Session, engine: "CaseEngine", inv: Invoice) -> None:
    """The lawyer bill is paid: the request is paid, the lawyer is put on the case (papers to sign) and both are
    told. The case stays unpaid for documents."""
    req = session.get(LawyerRequest, inv.lawyer_request_id) if inv.lawyer_request_id else None
    if req is None:
        return
    req.status = "paid"
    case = session.get(Case, req.case_id)
    app = session.get(LawyerApplication, req.application_id)
    if app is None:
        return
    if case.lawyer_application_id != app.id:
        assign(session, engine, case, app, "system:lawyer_paid")
    pack = engine.pack_of(case)
    lang = pack.lang(case.language)
    engine.notifier.notify(session, case, "lawyer", pack.t(
        lang, "notifications.lawyer_paid", name=lawyer_name(app),
        default=f"Оплата получена. Юрист {lawyer_name(app)} получил материалы дела. Подпишите соглашения с юристом "
                f"ЭЦП во вкладке «Юрист»."))
    lawyer = session.get(User, app.user_id) if app.user_id else None
    if lawyer is not None:
        engine.notifier.notify_user(session, lawyer, "lawyer", pack.t(
            lang, "notifications.lawyer_case_paid", default=(
                "Клиент оплатил вашу работу по делу. Досье дела открыто в кабинете юриста: "
                "https://konsilier.com/lawyer")))


def payout(inv: Invoice) -> Decimal:
    return money(Decimal(inv.amount) - Decimal(inv.commission_amount or 0))


def payment_view(session: Session, engine: "CaseEngine", inv: Invoice) -> dict[str, Any]:
    """What the client's payment view shows for a lawyer bill: the company's channel only."""
    view: dict[str, Any] = {"id": inv.id, "code": inv.code, "purpose": inv.purpose, "amount": float(inv.amount),
                            "currency": inv.currency, "status": inv.status}
    ch = channel(engine)
    if inv.status in engine.OPEN and ch is not None:
        view.update(ch)
    return view


def request_view(session: Session, engine: "CaseEngine", req: LawyerRequest) -> dict[str, Any]:
    app = session.get(LawyerApplication, req.application_id) if req.application_id else None
    inv = invoice_of(session, req)
    return {"id": req.id, "status": req.status, "application_id": req.application_id,
            "lawyer": {"name": lawyer_name(app), "kind": app.kind} if app else None,
            "price": float(req.price) if req.price is not None else None,
            "created_at": req.created_at.isoformat(),
            "invoice": payment_view(session, engine, inv) if inv is not None and inv.status != "cancelled" else None,
            "payment_available": channel(engine) is not None}
