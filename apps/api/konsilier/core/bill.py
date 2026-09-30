"""«Счёт на оплату» for a company or an individual entrepreneur paying from its bank account (PAYMENT_METHODS
bank_invoice). The form is not fixed by law; the bill carries what the payer's accountant needs to send the money:
the seller's name and tax number, bank, ИИК, БИК, Кбе, the КНП, the buyer, the item, the amount in figures and words, VAT, the
due date and a payment purpose with the bill number and payment code (the clients desk finds the payment by them).
Word file made with python-docx; the PDF comes from the same LibreOffice converter as the documents."""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

from .adapters.payment import Requisites

_ONES = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
_ONES_F = ["", "одна", "две", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
_TEENS = ["десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать", "пятнадцать", "шестнадцать",
          "семнадцать", "восемнадцать", "девятнадцать"]
_TENS = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят", "восемьдесят",
         "девяносто"]
_HUNDREDS = ["", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот", "восемьсот",
             "девятьсот"]
# (one, two-four, five+), feminine?
_SCALES = [(("", "", ""), False), (("тысяча", "тысячи", "тысяч"), True), (("миллион", "миллиона", "миллионов"), False),
           (("миллиард", "миллиарда", "миллиардов"), False)]


def _plural(n: int, forms: tuple[str, str, str]) -> str:
    n = abs(n) % 100
    if 11 <= n <= 19:
        return forms[2]
    n %= 10
    return forms[0] if n == 1 else forms[1] if 2 <= n <= 4 else forms[2]


def _triad(n: int, feminine: bool) -> list[str]:
    words = [_HUNDREDS[n // 100]]
    rest = n % 100
    if 10 <= rest <= 19:
        words.append(_TEENS[rest - 10])
    else:
        words.append(_TENS[rest // 10])
        words.append((_ONES_F if feminine else _ONES)[rest % 10])
    return [w for w in words if w]


def number_in_words(n: int) -> str:
    """0 ≤ n < 10¹²: «одна тысяча девятьсот девяносто»."""
    if n == 0:
        return "ноль"
    parts: list[str] = []
    for i, (forms, feminine) in enumerate(_SCALES):
        triad = (n // 1000 ** i) % 1000
        if not triad:
            continue
        words = _triad(triad, feminine)
        if i:
            words.append(_plural(triad, forms))
        parts = words + parts
    return " ".join(parts)


@dataclass
class BillWords:
    """Country words of the bill and the payment letter, from the pack's i18n «billing» block
    (packs/<cc>/i18n/ru.yaml): the tax-number labels, the currency in words and its sign."""
    seller_id: str = "ID"
    buyer_id: str = "ID"
    currency_forms: tuple[str, str, str] = ("", "", "")
    minor: str = ""
    sign: str = ""

    @classmethod
    def of(cls, pack: Any, currency: str | None) -> "BillWords":
        def t(key: str, default: str) -> str:
            return pack.t("ru", f"billing.{key}", default=default) if pack is not None else default

        forms = (t("currency_forms", currency or "").split("|") * 3)[:3]
        return cls(seller_id=t("seller_id", "ID"), buyer_id=t("buyer_id", "ID"),
                   currency_forms=(forms[0], forms[1], forms[2]), minor=t("minor", ""), sign=t("sign", currency or ""))


def amount_in_words(amount: Decimal, words: BillWords) -> str:
    """«Одна тысяча девятьсот девяносто тенге 00 тиын» (currency words from the pack)."""
    amount = Decimal(amount).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    whole, cents = int(amount), int((amount - int(amount)) * 100)
    text = f"{number_in_words(whole)} {_plural(whole, words.currency_forms)} {cents:02d} {words.minor}".strip()
    return text[0].upper() + text[1:]


def money(amount: Decimal) -> str:
    """«1 990,00»."""
    q = Decimal(amount).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    whole, frac = f"{q:.2f}".split(".")
    return f"{int(whole):,}".replace(",", " ") + "," + frac


ITEM_RU = {"document": "Подготовка юридического документа в сервисе Konsilier AI",
           "case": "«Дело под ключ»: подготовка документов по одному делу в сервисе Konsilier AI",
           "plan": "Доступ к сервису Konsilier AI по тарифу",
           "lawyer": "Оплата услуг юриста через сервис Konsilier AI"}


def bill_fields(req: Requisites, inv: Any, *, item: str, issued: date, words: BillWords) -> dict[str, Any]:
    """Everything the bill shows; one place for the Word file and the tests."""
    number = f"{inv.id}"
    amount = Decimal(inv.amount)
    vat = (amount * Decimal(12) / Decimal(112)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) if req.vat else None
    return {
        "number": number, "date": issued.strftime("%d.%m.%Y"),
        "due": (issued + timedelta(days=req.due_days)).strftime("%d.%m.%Y"),
        "seller": req.name, "seller_bin": req.bin, "seller_address": req.address, "bank": req.bank, "iik": req.iik,
        "bik": req.bik, "kbe": req.kbe, "knp": req.knp, "director": req.director,
        "buyer": inv.buyer_name or "", "buyer_bin": inv.buyer_bin or "", "buyer_address": inv.buyer_address or "",
        "item": item, "amount": money(amount), "words": amount_in_words(amount, words),
        "seller_id": words.seller_id, "buyer_id": words.buyer_id, "currency": words.currency_forms[2],
        "vat": f"в т. ч. НДС 12%: {money(vat)}" if vat is not None else "Без НДС",
        "purpose": f"Оплата по счёту № {number} от {issued.strftime('%d.%m.%Y')}, код {inv.code}. "
                   + (f"В т. ч. НДС {money(vat)}" if vat is not None else "Без НДС"),
        "code": inv.code,
    }


def make_bill_docx(fields: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    def row(table, cells: list[str], bold: bool = False) -> None:
        r = table.add_row().cells
        for c, text in zip(r, cells):
            c.text = text
            if bold:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.bold = True

    doc.add_paragraph("Образец платёжного поручения").runs[0].bold = True
    bank = doc.add_table(rows=0, cols=3)
    bank.style = "Table Grid"
    row(bank, ["Бенефициар:", "ИИК", "Кбе"], bold=True)
    row(bank, [f"{fields['seller']}\n{fields['seller_id']}: {fields['seller_bin']}", fields["iik"], fields["kbe"]])
    row(bank, ["Банк бенефициара:", "БИК", "Код назначения платежа"], bold=True)
    row(bank, [fields["bank"], fields["bik"], fields["knp"]])

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = h.add_run(f"Счёт на оплату № {fields['number']} от {fields['date']}")
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph(f"Поставщик: {fields['seller']}, {fields['seller_id']} {fields['seller_bin']}"
                      + (f", {fields['seller_address']}" if fields["seller_address"] else ""))
    doc.add_paragraph(f"Покупатель: {fields['buyer']}"
                      + (f", {fields['buyer_id']} {fields['buyer_bin']}" if fields["buyer_bin"] else "")
                      + (f", {fields['buyer_address']}" if fields["buyer_address"] else ""))

    items = doc.add_table(rows=0, cols=6)
    items.style = "Table Grid"
    row(items, ["№", "Наименование", "Кол-во", "Ед.", "Цена", "Сумма"], bold=True)
    row(items, ["1", fields["item"], "1", "усл.", fields["amount"], fields["amount"]])

    doc.add_paragraph(f"Итого: {fields['amount']} {fields['currency']}. {fields['vat']}.").runs[0].bold = True
    doc.add_paragraph(f"Всего к оплате: {fields['words']}.")
    doc.add_paragraph(f"Назначение платежа: {fields['purpose']}")
    doc.add_paragraph(f"Оплатить до {fields['due']}. Услуга оказывается после поступления оплаты. "
                      f"В назначении платежа укажите номер счёта и код {fields['code']}.")
    doc.add_paragraph("")
    doc.add_paragraph(f"Руководитель: {fields['director'] or '____________________'}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
