"""Manual filing on the official appeal portal (owner's decision 01.10.2026; docs/integrations-plan.md, 7.1 stage 1
and 7.3).

The person files the appeal on the portal themselves — in their own name, signing with their own digital signature.
We are only a technical aid: what to choose in the portal's form (pack data ``appeal_portal.yaml``), the ready
text and the PDF. Afterwards the person enters the appeal number and the date and may attach the portal's receipt;
we keep that as a ``Filing`` and run the response deadline and the reminders from the entered date.
"""

from __future__ import annotations

import hashlib
import io
import re
import uuid
from datetime import date, timedelta
from typing import Any

from docx import Document
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.concurrency import run_in_threadpool
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from ..container import Container
from ..core.appeal_portal import filing_record, portal_filing, portal_target, read_filing_proof
from ..core.engine import EngineError
from ..core.llm.base import Attachment
from ..core.models import Action, Case, Evidence, Filing, User, utcnow
from .deps import current_user, get_container, get_session, load_case
from .routes import _load_action, _read_upload, engine_error
from .views import case_view

router = APIRouter(prefix="/v1")

# letters (Latin, Cyrillic incl. national letters), digits and the separators portals use in registration numbers
NUMBER = re.compile(r"^[0-9A-Za-zА-Яа-яЁёӘәҒғҚқҢңӨөҰұҮүҺһІі№#\-/._ ]{3,64}$")
MAX_AGE_DAYS = 366  # a filing older than a year is not a «just filed» appeal: refuse a mistyped year


def _body_text(docx: bytes) -> str:
    """The document's text for the portal's text field: body paragraphs and tables, without page footers
    (the AI mark stays in the attached PDF)."""
    doc = Document(io.BytesIO(docx))
    parts = [p.text.strip() for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" · ".join(c.text.strip() for c in row.cells if c.text.strip()))
    text = "\n".join(parts)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _ready_action(container: Container, case: Case, action: Action) -> dict[str, Any]:
    if action.status not in ("ready", "submitted", "responded"):
        raise HTTPException(409, {"code": "awaiting_approval", "message": "document is awaiting lawyer approval"})
    if not container.engine.document_unlocked(case, action):
        raise HTTPException(402, {"code": "payment_required", "message": "document is available after payment"})
    target = portal_target(container.engine.pack_of(case), case.language, action)
    if target is None:
        raise HTTPException(409, {"code": "not_on_portal", "message": "this document is not filed on the appeal portal"})
    return target


@router.get("/cases/{case_id}/actions/{action_id}/portal-filing")
def portal_guide(case_id: uuid.UUID, action_id: uuid.UUID, user: User = Depends(current_user),
                   session: Session = Depends(get_session),
                   container: Container = Depends(get_container)) -> dict[str, Any]:
    """Step-by-step filing panel: body, category, ready text, the file to attach, and the record."""
    case = load_case(case_id, session, user)
    action = _load_action(case, action_id, session)
    target = _ready_action(container, case, action)
    text = _body_text(container.storage.get(action.docx_key)) if action.docx_key else ""
    f = portal_filing(session, action.id)
    return {**target, "text": text, "has_pdf": bool(action.pdf_key),
            "filing": filing_record(f) if f else None}


class PortalFiledIn(BaseModel):
    number: str = Field(min_length=1, max_length=80)
    filed_on: date
    receipt_evidence_id: uuid.UUID | None = None


def _receipt(session: Session, case: Case, evidence_id: uuid.UUID | None) -> Evidence | None:
    if evidence_id is None:
        return None
    ev = session.get(Evidence, evidence_id)
    if ev is None or ev.case_id != case.id:
        raise HTTPException(422, {"code": "bad_receipt", "message": "receipt not found in this case"})
    return ev


def _source(receipt: Evidence | None, number: str, filed_on: date) -> str:
    """``receipt_read`` when the person confirmed what we read from the confirmation unchanged, else
    ``client_entered`` (typed or corrected by hand)."""
    got = (receipt.extracted_facts or {}) if receipt is not None else {}
    if got.get("appeal_number") == number and got.get("filed_on") == filed_on.isoformat():
        return "receipt_read"
    return "client_entered"


@router.post("/cases/{case_id}/actions/{action_id}/portal-filing/proof", status_code=201)
async def portal_proof(case_id: uuid.UUID, action_id: uuid.UUID, file: UploadFile | None = File(default=None),
                       text: str | None = Form(default=None, max_length=20000),
                       user: User = Depends(current_user), session: Session = Depends(get_session),
                       container: Container = Depends(get_container)) -> dict[str, Any]:
    """«3 клика»: the person uploads the portal's confirmation (screenshot, PDF, the notification e-mail) or pastes
    its text; we keep it in the case as the receipt and read the appeal number and date from it, for a one-tap
    confirmation. Nothing is recorded as filed until the person confirms (``POST .../portal-filing``)."""
    case = load_case(case_id, session, user)
    action = _load_action(case, action_id, session)
    _ready_action(container, case, action)
    if file is not None:
        data, ctype = await _read_upload(file)
        name = file.filename or "confirmation"
    elif text and text.strip():
        data, ctype, name = text.strip().encode(), "text/plain", "confirmation.txt"
    else:
        raise HTTPException(422, {"code": "no_proof", "message": "attach the confirmation or paste its text"})
    engine = container.engine
    pack = engine.pack_of(case)

    def read() -> tuple[Evidence, dict[str, Any]]:  # a photo goes to the model: off the event loop
        ev = engine.add_evidence(session, case, kind="filing_receipt", filename=name, content_type=ctype, data=data)
        scan = ctype.startswith("image/") or (ctype == "application/pdf" and not ev.text)
        attachments = (Attachment(ctype, data, name),) if scan and engine.config.extract_images_with_llm else ()
        llm = engine.llm_for(case)
        got = read_filing_proof(llm, pack.appeal_portal, pack.lang(case.language), ev.text, attachments,
                                pack.local_now().date())
        engine._save_vault(case, llm)
        return ev, got
    ev, got = await run_in_threadpool(read)
    number, filed_on = got["number"], got["filed_on"]
    ev.extracted_facts = {"appeal_number": number, "filed_on": filed_on.isoformat() if filed_on else None,
                          "read_by": got["source"]}
    ev.confirmed = False
    engine.audit(session, case, f"user:{user.id}", "filing_proof_read", action=action.action_id,
                 found=bool(number and filed_on), read_by=got["source"])
    session.flush()
    return {"evidence_id": str(ev.id), "number": number, "filed_on": filed_on.isoformat() if filed_on else None,
            "found": bool(number and filed_on), "read_by": got["source"]}


@router.post("/cases/{case_id}/actions/{action_id}/portal-filing", status_code=201)
def portal_filed(case_id: uuid.UUID, action_id: uuid.UUID, body: PortalFiledIn,
                   user: User = Depends(current_user), session: Session = Depends(get_session),
                   container: Container = Depends(get_container)) -> dict[str, Any]:
    """The person filed the appeal on the portal and entered its number and date: keep the proof, start the
    response deadline from that date."""
    case = load_case(case_id, session, user)
    action = _load_action(case, action_id, session)
    target = _ready_action(container, case, action)
    number = " ".join(body.number.split())
    if not NUMBER.match(number) or not re.search(r"\d", number):
        raise HTTPException(422, {"code": "bad_number", "message": "enter the appeal number from the portal"})
    pack = container.engine.pack_of(case)
    today = pack.local_now().date()
    if body.filed_on > today:
        raise HTTPException(422, {"code": "date_in_future", "message": "the filing date is in the future"})
    made = action.created_at.astimezone(pack.tz).date() if action.created_at and action.created_at.tzinfo \
        else (action.created_at.date() if action.created_at else today)
    if body.filed_on < made or body.filed_on < today - timedelta(days=MAX_AGE_DAYS):
        raise HTTPException(422, {"code": "date_before_document",
                                  "message": "the filing date is earlier than the document"})
    receipt = _receipt(session, case, body.receipt_evidence_id)
    if portal_filing(session, action.id) is not None:
        raise HTTPException(409, {"code": "already_filed", "message": "this document is already marked as filed"})
    key, fmt = (action.pdf_key, "pdf") if action.pdf_key else (action.docx_key, "docx")
    sha = hashlib.sha256(container.storage.get(key)).hexdigest() if key else None
    try:
        container.engine.mark_submitted(session, case, action, f"user:{user.id}", via=target["channel"],
                                        submitted_on=body.filed_on)
    except EngineError as e:
        raise engine_error(e) from e
    source = _source(receipt, number, body.filed_on)
    f = Filing(case_id=case.id, action_id=action.id, user_id=user.id, channel=target["channel"],
               recipient=(target["recipient"] or target["body"] or "")[:500],
               body=(target["body"] or target["recipient"] or "")[:500], body_key=target["body_key"],
               appeal_type=target["appeal_type"], category=(target["category"] or None),
               appeal_number=number, filed_at=body.filed_on, status="filed",
               receipt_evidence_id=receipt.id if receipt else None,
               doc_format=fmt if key else None, doc_sha256=sha, source=source, attachments=[],
               events=[{"at": utcnow().isoformat(), "type": "filed", "source": source,
                        "filed_at": body.filed_on.isoformat()}])
    session.add(f)
    container.engine.audit(session, case, f"user:{user.id}", "filing_recorded", action=action.action_id,
                           channel=target["channel"], number=number, filed_at=body.filed_on.isoformat(),
                           receipt=bool(receipt), doc_sha256=sha)
    session.flush()
    return {"filing": filing_record(f), "case": case_view(container.engine, session, case)}


class ReceiptIn(BaseModel):
    evidence_id: uuid.UUID


@router.post("/cases/{case_id}/actions/{action_id}/portal-filing/receipt")
def portal_receipt(case_id: uuid.UUID, action_id: uuid.UUID, body: ReceiptIn,
                     user: User = Depends(current_user), session: Session = Depends(get_session),
                     container: Container = Depends(get_container)) -> dict[str, Any]:
    """Attach the portal's receipt or a screenshot (already uploaded as evidence) to a filing made earlier."""
    case = load_case(case_id, session, user)
    action = _load_action(case, action_id, session)
    f = portal_filing(session, action.id)
    if f is None:
        raise HTTPException(404, {"code": "no_filing", "message": "no filing for this document"})
    receipt = _receipt(session, case, body.evidence_id)
    f.receipt_evidence_id = receipt.id if receipt else None
    f.events = [*(f.events or []), {"at": utcnow().isoformat(), "type": "receipt", "source": "client"}]
    container.engine.audit(session, case, f"user:{user.id}", "filing_receipt_added", action=action.action_id)
    session.flush()
    return {"filing": filing_record(f), "case": case_view(container.engine, session, case)}
