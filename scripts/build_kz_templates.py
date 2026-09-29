"""Generate the DOCX templates of the KZ pack (docxtpl / Jinja2 syntax).

Templates are pack DATA: a lawyer can open and edit them in Word directly.
This script only bootstraps the first versions:

    python scripts/build_kz_templates.py

Available context (see core/engine.py::document_context):
    title, f.<field>, labels.<field>, applicant.{name,id,address,email,kind},
    respondent.{name,id,address,email,kind}, addressee.{name,id,address,email,key,kind},
    narrative, demands, norm_refs, evidence, previous_actions, date, currency
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1] / "packs" / "kz" / "templates"

NORMS = ("{% for r in norm_refs %}{{ '[норма уточняется юристом]' if 'TODO' in r else r }}"
         "{% if not loop.last %}; {% endif %}{% endfor %}")


def _doc() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    return doc


def _right(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(text)


def _center(doc: Document, text: str, bold: bool = True) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold


def _para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.bold = bold


def _opt(doc: Document, cond: str, text: str, right: bool = False) -> None:
    """A paragraph printed only when ``cond`` is true ({%p %} removes the whole paragraph otherwise)."""
    doc.add_paragraph("{%p if " + cond + " %}")
    if right:
        _right(doc, text)
    else:
        _para(doc, text)
    doc.add_paragraph("{%p endif %}")


def _header(doc: Document, *, birth_date: bool = False) -> None:
    """Addressee and applicant requisites.

    The applicant block follows АППК РК ст. 63 п. 2 / ст. 93 п. 2 (ФИО, ИИН, почтовый адрес) and, for a lawsuit,
    ГПК РК ст. 148 ч. 2 пп. 2) (дата рождения, место жительства, ИИН, телефон и e-mail, если есть)."""
    _right(doc, "Кому: {{ addressee.name }}")
    _opt(doc, "addressee.id", "{{ 'ИИН' if addressee.kind == 'person' else 'БИН' }}: {{ addressee.id }}", right=True)
    _opt(doc, "addressee.address", "Адрес: {{ addressee.address }}", right=True)
    _opt(doc, "addressee.email", "E-mail: {{ addressee.email }}", right=True)
    _right(doc, "От: {{ applicant.name }}")
    if birth_date:
        _opt(doc, "f.applicant_birth_date", "Дата рождения: {{ f.applicant_birth_date }}", right=True)
    _opt(doc, "applicant.id", "ИИН: {{ applicant.id }}", right=True)
    _opt(doc, "applicant.address", "Адрес: {{ applicant.address }}", right=True)
    _right(doc, "Тел.: {{ f.applicant_phone }}")
    _opt(doc, "f.applicant_email", "E-mail: {{ f.applicant_email }}", right=True)


# who the complaint / statement is about, when the document goes to a body and not to that party itself
RESPONDENT = ("{{ respondent.name }}{% if respondent.id %}, {{ 'ИИН' if respondent.kind == 'person' else 'БИН' }} "
              "{{ respondent.id }}{% endif %}{% if respondent.address %}, адрес: {{ respondent.address }}{% endif %}"
              "{% if respondent.email %}, e-mail: {{ respondent.email }}{% endif %}")


def _about(doc: Document, label: str) -> None:
    _opt(doc, "respondent.name and addressee.key != 'respondent'", label + ": " + RESPONDENT + ".")


def _facts(doc: Document) -> None:
    """Dates and amounts under the scenario's own labels (e.g. «Дата ДТП», «Сумма штрафа»)."""
    _opt(doc, "f.event_date or f.amount or f.decision_number",
         "{% if f.decision_number %}{{ labels.decision_number }}: {{ f.decision_number }}. {% endif %}"
         "{% if f.event_date %}{{ labels.event_date }}: {{ f.event_date }}.{% endif %}"
         "{% if f.amount %} {{ labels.amount }}: {{ f.amount }} {{ currency }}.{% endif %}")


def _previous(doc: Document, title: str = "Ранее предпринятые действия:") -> None:
    doc.add_paragraph("{%p if previous_actions %}")
    _para(doc, title, bold=True)
    doc.add_paragraph("{%p for a in previous_actions %}")
    doc.add_paragraph("— {{ a.title }} (направлено {{ a.date }}): {{ a.response }}.")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p endif %}")


def _attachments(doc: Document, extra: tuple[str, ...] = ()) -> None:
    """Перечень приложений: files the person added, copies of the earlier documents and their answers
    (Закон о ЗПП ст. 42-5 п. 3; АППК ст. 93 п. 2 пп. 8); ``extra`` — items the law requires for this kind."""
    doc.add_paragraph("{%p if evidence or previous_actions" + (" or true" if extra else "") + " %}")
    _para(doc, "Приложения:", bold=True)
    doc.add_paragraph("{%p set n = namespace(i=0) %}")
    doc.add_paragraph("{%p for e in evidence %}")
    doc.add_paragraph("{% set n.i = n.i + 1 %}{{ n.i }}. {{ e }}")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p for a in previous_actions %}")
    doc.add_paragraph("{% set n.i = n.i + 1 %}{{ n.i }}. Копия документа «{{ a.title }}» от {{ a.date }}"
                      "{% if a.response %} и ответ на него (или сведения о его отсутствии){% endif %}.")
    doc.add_paragraph("{%p endfor %}")
    for item in extra:
        doc.add_paragraph("{% set n.i = n.i + 1 %}{{ n.i }}. " + item)
    doc.add_paragraph("{%p endif %}")


def _sign(doc: Document) -> None:
    doc.add_paragraph("")
    doc.add_paragraph("Дата: {{ date }}                    Подпись: ______________ / {{ applicant.name }}")


def _tail(doc: Document, demand_intro: str, norms: str, extra_attachments: tuple[str, ...] = ()) -> None:
    _para(doc, demand_intro, bold=True)
    _para(doc, "{{ demands }}")
    _para(doc, "Правовое основание: " + norms)
    _attachments(doc, extra_attachments)
    _sign(doc)


def claim_seller() -> Document:
    doc = _doc()
    _header(doc)
    _center(doc, "ПРЕТЕНЗИЯ")
    _center(doc, "о возврате уплаченной денежной суммы", bold=False)
    _para(doc, "{{ f.purchase_date }} мной приобретен(а) «{{ f.goods_description }}» у {{ addressee.name }} "
               "на сумму {{ f.amount }} {{ currency }}.")
    _para(doc, "{{ narrative }}")
    _tail(doc, "На основании изложенного прошу:", NORMS)
    return doc


def complaint_authority() -> Document:
    """Обращение потребителя в госорган: Закон о ЗПП ст. 42-5 п. 4 (данные потребителя, данные продавца —
    наименование, почтовый адрес, БИН; требование и обстоятельства), п. 3 (копия претензии и ответа)."""
    doc = _doc()
    _header(doc)
    _center(doc, "ЖАЛОБА")
    _center(doc, "о нарушении прав потребителя", bold=False)
    _para(doc, "Продавец (исполнитель): " + RESPONDENT + ".")
    _para(doc, "{{ f.purchase_date }} мной приобретен(а) «{{ f.goods_description }}» у {{ respondent.name }} "
               "на сумму {{ f.amount }} {{ currency }}.")
    _para(doc, "{{ narrative }}")
    _previous(doc)
    _tail(doc, "Прошу:", NORMS)
    return doc


def statement_lender() -> Document:
    doc = _doc()
    _header(doc)
    _center(doc, "ЗАЯВЛЕНИЕ")
    _center(doc, "об оформлении займа без волеизъявления заемщика (мошенничество)", bold=False)
    _para(doc, "Мне стало известно, что {{ f.loan_date }} на мое имя в {{ addressee.name }} оформлен заем "
               "на сумму {{ f.amount }} {{ currency }}{% if f.contract_number %} (договор № {{ f.contract_number }})"
               "{% endif %}. Данный заем я не оформлял(а), согласия на его оформление не давал(а), денежные "
               "средства не получал(а) и ими не распоряжался(ась).")
    _para(doc, "{{ narrative }}")
    _opt(doc, "f.police_report_number",
         "По факту мошенничества я обратился(ась) в органы полиции: {{ f.police_report_number }}.")
    _tail(doc, "Прошу:", NORMS)
    return doc


def complaint_arrf() -> Document:
    doc = _doc()
    _header(doc)
    _center(doc, "ЖАЛОБА")
    _center(doc, "на действия финансовой организации", bold=False)
    _para(doc, "Финансовая организация: " + RESPONDENT + ".")
    _para(doc, "{{ f.loan_date }} на мое имя в {{ respondent.name }} без моего волеизъявления оформлен заем на сумму "
               "{{ f.amount }} {{ currency }}{% if f.contract_number %} (договор № {{ f.contract_number }}){% endif %}.")
    _para(doc, "{{ narrative }}")
    _opt(doc, "f.police_report_number", "Обращение в полицию: {{ f.police_report_number }}.")
    _previous(doc)
    _tail(doc, "Прошу:", NORMS)
    return doc


GENERIC_NORMS = ("{% for r in norm_refs %}{{ '[норма: уточнит юрист]' if 'TODO' in r else r }}"
                 "{% if not loop.last %}; {% endif %}{% endfor %}")


def generic(title: str, subtitle: str, demand_intro: str, about: str) -> Document:
    """Universal document (statement, complaint, appeal, motion): norms/deadlines only from pack data."""
    doc = _doc()
    _header(doc)
    _center(doc, title)
    _center(doc, subtitle, bold=False)
    _about(doc, about)
    _facts(doc)
    _para(doc, "{{ narrative }}")
    _previous(doc)
    _tail(doc, demand_intro, GENERIC_NORMS)
    return doc


def lawsuit() -> Document:
    """Исковое заявление — ГПК РК ст. 148 ч. 2: суд; истец (ФИО, дата рождения, место жительства, ИИН, телефон,
    e-mail); ответчик (ФИО/наименование, место жительства/нахождения, ИИН/БИН, если известны); суть нарушения и
    требования; обстоятельства и доказательства; досудебный порядок; цена иска; перечень приложений.
    Приложения — ГПК ст. 149 ч. 1: копии по числу ответчиков, документ об уплате госпошлины."""
    doc = _doc()
    _header(doc, birth_date=True)
    _right(doc, "Ответчик: {{ respondent.name }}")
    _opt(doc, "respondent.id", "{{ 'ИИН' if respondent.kind == 'person' else 'БИН' }}: {{ respondent.id }}", right=True)
    _right(doc, "{{ 'Место нахождения' if respondent.kind == 'business' else 'Место жительства' }}: "
                "{{ respondent.address or '[не известно истцу — уточнит юрист]' }}")
    _opt(doc, "respondent.email", "E-mail: {{ respondent.email }}", right=True)
    _right(doc, "Цена иска: {% if f.amount %}{{ f.amount }} {{ currency }}{% else %}[рассчитает юрист]{% endif %}")
    _center(doc, "ИСКОВОЕ ЗАЯВЛЕНИЕ")
    _center(doc, "{{ title }}", bold=False)
    _facts(doc)
    _opt(doc, "f.children_info", "{{ labels.children_info }}: {{ f.children_info }}.")
    _para(doc, "{{ narrative }}")
    _previous(doc, "Досудебное обращение к ответчику:")
    _para(doc, "На основании изложенного прошу суд:", bold=True)
    _para(doc, "{{ demands }}")
    _para(doc, "Правовое основание: " + GENERIC_NORMS)
    _attachments(doc, ("Копия искового заявления и приложенных к нему документов по числу ответчиков.",
                       "Документ, подтверждающий уплату государственной пошлины, если она уплачивается "
                       "(размер и основания освобождения уточнит юрист)."))
    _sign(doc)
    return doc


def claim_letter() -> Document:
    """Universal pre-trial claim to the other party itself (seller, bank, employer, landlord, debtor)."""
    doc = _doc()
    _header(doc)
    _center(doc, "ПРЕТЕНЗИЯ")
    _center(doc, "(досудебная)", bold=False)
    _facts(doc)
    _para(doc, "{{ narrative }}")
    _previous(doc)
    _para(doc, "На основании изложенного требую:", bold=True)
    _para(doc, "{{ demands }}")
    _para(doc, "Правовое основание: " + GENERIC_NORMS)
    _para(doc, "Прошу дать письменный ответ на претензию. Если требования не будут удовлетворены, я буду вынужден(а) "
               "обратиться в уполномоченный государственный орган и (или) в суд для защиты своих прав.")
    _attachments(doc)
    _sign(doc)
    return doc


def main() -> None:
    out = {
        "consumer/claim_seller.docx": claim_seller,
        "consumer/complaint_authority.docx": complaint_authority,
        "money/statement_lender.docx": statement_lender,
        "money/complaint_arrf.docx": complaint_arrf,
        "generic/complaint.docx": lambda: generic(
            "ЖАЛОБА", "{{ title }}", "Прошу:", "Лицо (орган), на действия (бездействие) которого подается жалоба"),
        "generic/statement.docx": lambda: generic(
            "ЗАЯВЛЕНИЕ", "{{ title }}", "Прошу:", "Лицо, о действиях которого подается заявление"),
        "generic/lawsuit.docx": lawsuit,
        "generic/appeal.docx": lambda: generic("АПЕЛЛЯЦИОННАЯ ЖАЛОБА", "{{ title }}", "Прошу:", "Ответчик"),
        "generic/claim_letter.docx": claim_letter,
        "generic/motion.docx": lambda: generic(
            "ХОДАТАЙСТВО", "{{ title }}{% if f.case_number %} (дело № {{ f.case_number }}){% endif %}", "Прошу:",
            "Лицо, о действиях которого подается ходатайство"),
    }
    only = set(sys.argv[1:])  # e.g. generic/claim_letter.docx — rebuild just these
    for rel, build in out.items():
        if only and rel not in only:
            continue
        path = ROOT / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        build().save(path)
        print("wrote", path)


if __name__ == "__main__":
    main()
