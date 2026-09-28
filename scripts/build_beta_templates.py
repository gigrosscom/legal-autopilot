"""Generate the DOCX templates of the KZ beta scenarios (services and business), docxtpl / Jinja2 syntax.

    python scripts/build_beta_templates.py

services/package.docx — a package of documents of a service scenario (tender bid, admission, visa, ИП):
    L.* (fixed words in the case language), scenario_title, disclaimer, addressee.name, applicant.*,
    sections[] {title, paragraphs[], check[], signature}, sources[] {title, url, checked_on}, date
business/claim_b2b.docx, business/complaint_b2b.docx — like the generic claim/complaint, but for a business
    applicant (ИП / ТОО): БИН/ИИН, representative, contract number and date; after the demands, the paragraphs of the
    action's optional ``package`` sections (e.g. a contractual penalty or the legal basis in acts that are not key acts).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1] / "packs" / "kz" / "templates"

NORMS = ("{% for r in norm_refs %}{{ '[норма: уточнит юрист]' if 'TODO' in r else r }}"
         "{% if not loop.last %}; {% endif %}{% endfor %}")


def _doc() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    return doc


def _p(doc: Document, text: str, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold: bool = False, italic: bool = False,
       size: int | None = None, page_break: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.page_break_before = page_break
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)


def _tag(doc: Document, tag: str) -> None:
    doc.add_paragraph(tag)


R, C = WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER


def package() -> Document:
    doc = _doc()
    _p(doc, "{{ L.to }}: {{ addressee.name }}", align=R)
    _p(doc, "{{ L.from }}: {{ applicant.name }}", align=R)
    _p(doc, "{% if applicant.id %}{{ id_label }}: {{ applicant.id }}{% endif %}", align=R)
    _p(doc, "{% if f.applicant_phone %}{{ L.phone }}: {{ f.applicant_phone }}{% endif %}", align=R)
    _p(doc, "{% if f.applicant_email %}E-mail: {{ f.applicant_email }}{% endif %}", align=R)
    _p(doc, "{{ L.package|upper }}", align=C, bold=True)
    _p(doc, "{{ scenario_title }}", align=C)
    _p(doc, "{{ disclaimer }}", italic=True, size=10)
    _p(doc, "{{ L.contents }}:", bold=True)
    _tag(doc, "{%p for s in sections %}")
    _p(doc, "{{ loop.index }}. {{ s.title }}")
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p for s in sections %}")
    _p(doc, "{{ s.title }}", align=C, bold=True, page_break=True)
    _tag(doc, "{%p for x in s.paragraphs %}")
    _p(doc, "{{ x }}")
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p for x in s.check %}")
    _p(doc, "☐ {{ x }}", align=WD_ALIGN_PARAGRAPH.LEFT)
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p if s.signature %}")
    _p(doc, "")
    _p(doc, "{{ L.date }}: {{ date }}                    {{ L.signature }}: ______________ / {{ applicant.name }}")
    _tag(doc, "{%p endif %}")
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p if sources %}")
    _p(doc, "{{ L.sources }}", align=C, bold=True, page_break=True)
    _tag(doc, "{%p for s in sources %}")
    _p(doc, "{{ s.title }} — {{ s.url }}{% if s.checked_on %} ({{ L.checked }} {{ s.checked_on }}){% endif %}",
       align=WD_ALIGN_PARAGRAPH.LEFT, size=10)
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p endif %}")
    return doc


def _b2b_header(doc: Document) -> None:
    _p(doc, "Кому: {{ addressee.name }}", align=R)
    _p(doc, "{% if addressee.id %}БИН/ИИН: {{ addressee.id }}{% endif %}", align=R)
    _p(doc, "{% if addressee.address %}Адрес: {{ addressee.address }}{% endif %}", align=R)
    _p(doc, "От: {{ applicant.name }}", align=R)
    _p(doc, "{% if applicant.id %}БИН/ИИН: {{ applicant.id }}{% endif %}", align=R)
    _p(doc, "{% if f.representative %}в лице: {{ f.representative }}{% endif %}", align=R)
    _p(doc, "Адрес: {{ f.applicant_address }}", align=R)
    _p(doc, "Тел.: {{ f.applicant_phone }}", align=R)


def _b2b_body(doc: Document) -> None:
    _p(doc, "{% if f.contract_number %}Договор № {{ f.contract_number }}{% if f.contract_date %} от "
            "{{ f.contract_date }}{% endif %}.{% endif %}{% if f.event_date %} Дата события: {{ f.event_date }}."
            "{% endif %}{% if f.amount %} Сумма требований: {{ f.amount }} {{ currency }}.{% endif %}")
    _p(doc, "{{ narrative }}")
    _tag(doc, "{%p if previous_actions %}")
    _p(doc, "Ранее предпринятые действия:", bold=True)
    _tag(doc, "{%p for a in previous_actions %}")
    _p(doc, "— {{ a.title }} (направлено {{ a.date }}): {{ a.response }}.")
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p endif %}")


def _b2b_tail(doc: Document, intro: str, closing: str) -> None:
    _p(doc, intro, bold=True)
    _p(doc, "{{ demands }}")
    _tag(doc, "{%p for s in sections %}")
    _tag(doc, "{%p for x in s.paragraphs %}")
    _p(doc, "{{ x }}")
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p endfor %}")
    _p(doc, "{% if norm_refs %}Правовое основание: " + NORMS + "{% endif %}")
    if closing:
        _p(doc, closing)
    _tag(doc, "{%p if evidence %}")
    _p(doc, "Приложения:", bold=True)
    _tag(doc, "{%p for e in evidence %}")
    _p(doc, "{{ loop.index }}. {{ e }}")
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p endif %}")
    _p(doc, "")
    _p(doc, "Дата: {{ date }}          Подпись: ______________ / {{ f.representative or applicant.name }}"
            "          М.П. (при наличии)")


def claim_b2b() -> Document:
    doc = _doc()
    _b2b_header(doc)
    _p(doc, "ПРЕТЕНЗИЯ", align=C, bold=True)
    _p(doc, "{{ title }}", align=C)
    _b2b_body(doc)
    _b2b_tail(doc, "На основании изложенного требуем:",
              "Просим дать письменный ответ на претензию. При неудовлетворении требований мы оставляем за собой "
              "право обратиться в суд.")
    return doc


def complaint_b2b() -> Document:
    doc = _doc()
    _b2b_header(doc)
    _p(doc, "ЖАЛОБА", align=C, bold=True)
    _p(doc, "{{ title }}", align=C)
    _p(doc, "Лицо (орган), на действия (решение) которого подается жалоба: {{ f.respondent_name }}.")
    _b2b_body(doc)
    _b2b_tail(doc, "Прошу:", "")
    return doc


def main() -> None:
    for rel, build in (("services/package.docx", package), ("business/claim_b2b.docx", claim_b2b),
                       ("business/complaint_b2b.docx", complaint_b2b)):
        path = ROOT / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        build().save(path)
        print("wrote", path)


if __name__ == "__main__":
    main()
